"""Add the completed sharing experiment without changing the original framework."""
from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Inches,Pt
OUT=Path('outputs/sharing_study');a=json.loads((OUT/'analysis.json').read_text());narrative=json.loads((OUT/'interpretation.json').read_text())
d=Document('outputs/architecture_revision/E2E_PREF_architecture_focused.docx')
def replace(p,t):
 props=deepcopy(p.runs[0]._r.rPr) if p.runs else None;p.clear();r=p.add_run(t)
 if props is not None:r._r.insert(0,props)
def find(prefix):return next(p for p in d.paragraphs if p.text.startswith(prefix))
def before(anchor,text,template):
 p=anchor.insert_paragraph_before(text)
 if template._p.pPr is not None:p._p.insert(0,deepcopy(template._p.pPr))
 if text and template.runs and template.runs[0]._r.rPr is not None:p.runs[0]._r.insert(0,deepcopy(template.runs[0]._r.rPr))
 return p
body=find('The experiments examine');sub=find('5.2 Contributions');caption=find('Fig. 3. Component')
replace(body,'The experiments examine complete-framework ranking across domains, the consequences of sharing graph representations between pathways, component substitutions, and the effects of auxiliary learning or continued encoder updating. A fresh MovieLens-20M cohort [4] and an independent product domain provide complementary evidence about these choices.')
anchor=find('Candidate-conditioned interest modeling')
before(anchor,'TGSRec [30] integrates temporal collaborative signals and sequential patterns through a continuous-time user–item graph and a temporal collaborative transformer. Thus, combining graph and sequential learning is an established direction. E2E-PREF uses a static item co-occurrence graph whose representations feed both a causal history encoder and item-content fusion. Its sharing experiment evaluates this specific connection against separate graph pathways.',body)
anchor=find('2.4 Baselines')
before(anchor,'IAFCL [31] combines item–attribute graph fusion with contrastive sequential recommendation. Its attribute aggregation differs from E2E-PREF’s use of a co-occurrence graph shared between the history and item pathways. Together with TGSRec, this work motivates evaluating the proposed connection explicitly.',body)
p=find('BPR formulates implicit-feedback')
replace(p,p.text+' Sharing item embeddings between sequence input and prediction also appears in SASRec [3]; E2E-PREF evaluates a specific graph-sharing connection across history and content pathways.')
p=find('The design differs in its coupling')
p._p.getparent().remove(p._p)
# Renumber existing results before inserting the new figure.
for p in d.paragraphs:
 t=p.text
 if t.startswith('5.2 Contributions'):t=t.replace('5.2','5.3',1)
 elif t.startswith('5.3 Auxiliary'):t=t.replace('5.3','5.4',1)
 elif t.startswith('5.4 Adaptation'):t=t.replace('5.4','5.5',1)
 t=re.sub(r'(Fig\. |Figure )([345])\b',lambda m:m[1]+str(int(m[2])+1),t)
 if t!=p.text:replace(p,t)
anchor=find('5 Findings')
before(anchor,'A follow-up sharing experiment compares the original model with two independent graph encoders, one for history and one for item fusion. A control uses the same two encoders but feeds their mean to both pathways, matching the separate model’s parameter count. Both variants receive the same three-rate search and five-seed confirmation protocol. All additional fitting completes before new test evaluation. This is a post hoc extension on the existing datasets; the local analysis plan was fixed before the new runs.',body)
anchor=find('5.3 Contributions')
before(anchor,'5.2 Shared Representations Across Pathways',sub)
before(anchor,narrative['design'],body)
p=before(anchor,'',body);p.alignment=1;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(OUT/'figure3_sharing.png'),width=Inches(6.85))
p=before(anchor,'Fig. 3. Direct representation-sharing comparisons. Points show paired full-catalog NDCG@10 differences and lines show 95% conditional seed/user bootstrap intervals. Positive values favor the shared design. The original model uses fewer parameters than separate encoders; the shared-average control and separate encoders have identical parameter counts. All comparisons are exploratory.',caption)
for r in p.runs:r.font.size=Pt(9)
p.paragraph_format.keep_together=True
before(anchor,narrative['results'],body)
anchor=find('Candidate-aware scoring also connects')
before(anchor,narrative['discussion'],body)
replace(find('The evidence covers one fresh'), 'The evidence covers one fresh MovieLens cohort and one small Amazon product domain under static catalog information. The sharing study is a post hoc extension on these same datasets and its equal-parameter control changes both routing and averaging. The evidence does not establish universal superiority or performance against untested integrated models such as TGSRec and IAFCL. Future work will examine additional domains, direct integrated-model comparisons, broader tuning, alternative graph and temporal configurations, and retrieval and online evaluation. Detailed timestamp and uncertainty limitations remain in the supplement.')
replace(find('E2E-PREF connects collaborative, temporal and content representations through'),narrative['conclusion'])
# Preserve abstract's accessible level; add no implementation detail.
p=find('Abstract—');replace(p,p.text.replace('Component analyses identify domain-dependent roles for auxiliary learning and preference signals.','Direct architectural comparisons characterize representation sharing, while component analyses identify domain-dependent roles for auxiliary learning and preference signals.'))
p=find('Each learned configuration receives')
replace(p,p.text.replace('Validation selects the rate and checkpoint','Sampled-validation NDCG@10 selects the rate and checkpoint'))
# Add verified references then renumber all citations by first appearance.
ref_template=find('[29] ')
for text in ['[30] Z. Fan, Z. Liu, J. Zhang, Y. Xiong, L. Zheng, and P. S. Yu, “Continuous-time sequential recommendation with temporal graph collaborative transformer,” in Proc. CIKM, 2021. doi: 10.1145/3459637.3482242.', '[31] D. Zhang, J. Qin, J. Ma, Z. Yang, D. Cui, and P. Ji, “Item attributes fusion based on contrastive learning for sequential recommendation,” Multimedia Systems, vol. 30, Art. 291, 2024. doi: 10.1007/s00530-024-01486-7.']:
 p=d.add_paragraph(text)
 if ref_template._p.pPr is not None:p._p.insert(0,deepcopy(ref_template._p.pPr))
 if ref_template.runs[0]._r.rPr is not None:p.runs[0]._r.insert(0,deepcopy(ref_template.runs[0]._r.rPr))
refs=[p for p in d.paragraphs if re.match(r'^\[\d+\] ',p.text)];reftext={int(re.match(r'^\[(\d+)\]',p.text)[1]):p.text for p in refs};mapping={}
for p in d.paragraphs:
 if p.text=='References':break
 for x in re.findall(r'\[(\d+)\]',p.text):
  if int(x) not in mapping:mapping[int(x)]=len(mapping)+1
assert len(mapping)==31
renumber=lambda t:re.sub(r'\[(\d+)\]',lambda m:'['+str(mapping[int(m[1])])+']',t)
for p in d.paragraphs:
 if p.text=='References':break
 if re.search(r'\[\d+\]',p.text):replace(p,renumber(p.text))
for p,(old,new) in zip(refs,sorted(mapping.items(),key=lambda kv:kv[1])):replace(p,renumber(reftext[old]))
assert len(d.tables)==0 and len(d.inline_shapes)==6
assert len(d._element.xpath('.//m:oMath'))==3
for p in d.paragraphs:
 if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
path=OUT/'E2E_PREF_framework_evaluated.docx';d.save(path)
# Supplement preserves full old technical record with updated reference numbering.
s=Path('outputs/architecture_revision/Technical_supplement.md').read_text();start=s.index('\n3 The E2E-PREF Framework');end=s.index('\nReferences',start);eq=s[s.index('## Linear equation transcriptions'):]
old=renumber(s[start:end]);references='\n\n'.join(p.text for p in refs)
new='# E2E-PREF technical supplement\n\nSupporting material for the framework manuscript including the representation-sharing extension.\n\n## S1 Representation-sharing experiment\n\n'+Path('work/sharing_study/PROTOCOL.md').read_text().split('\n',1)[1]+'\n\n'+(OUT/'RESULTS.md').read_text()+'\n\n### Interpretation\n\n'+narrative['results']+'\n\n'+narrative['discussion']+'\n\n## S2 Original implementation and numerical records\n\nSection and table numbers below retain their technical-record numbering. Reference numbers match the revised main manuscript.\n'+old+'\n\n## References\n\n'+references+'\n\n'+eq
(OUT/'Technical_supplement.md').write_text(new)
(OUT/'reference_number_mapping.json').write_text(json.dumps(mapping,indent=2))
print(path)
